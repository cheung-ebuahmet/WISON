#!/usr/bin/env python3
"""
Generate a Professional Work Handover Memo (.docx)
For temporary leave / absence handover

Covers: system infrastructure, project knowledge base, ongoing work,
recurring tasks, key contacts, and urgent items.
"""

import datetime
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

today = datetime.date.today()
today_str = today.strftime("%d-%b-%Y")

# ── Colours ──────────────────────────────────────────────────────────────────
DARK_BLUE   = RGBColor(0x0B, 0x2A, 0x4A)
MED_BLUE    = RGBColor(0x1F, 0x5C, 0x8A)
RED_URGENT  = RGBColor(0xC0, 0x39, 0x2B)
ACCENT_GREY = RGBColor(0x6B, 0x7B, 0x8D)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GREY   = RGBColor(0x33, 0x3F, 0x4D)
BORDER_GREY = RGBColor(0xBF, 0xCB, 0xD7)

# ── Helpers ──────────────────────────────────────────────────────────────────
doc = Document()

section = doc.sections[0]
section.page_width  = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin    = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin   = Cm(2.5)
section.right_margin  = Cm(2.0)

style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(10)
font.color.rgb = DARK_GREY

def add_table(doc, headers, rows, col_widths=None):
    """Professional table with dark header."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True; run.font.size = Pt(9); run.font.color.rgb = WHITE; run.font.name = 'Calibri'
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0B2A4A"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
    for r, row_data in enumerate(rows):
        row = table.rows[r + 1]
        for c, text in enumerate(row_data):
            cell = row.cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(text))
            run.font.size = Pt(9); run.font.name = 'Calibri'
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            if r % 2 == 1:
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F6FA"/>')
                cell._tc.get_or_add_tcPr().append(shading)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:color="BFCBD7"/>'
        f'<w:left w:val="single" w:sz="4" w:color="BFCBD7"/>'
        f'<w:bottom w:val="single" w:sz="4" w:color="BFCBD7"/>'
        f'<w:right w:val="single" w:sz="4" w:color="BFCBD7"/>'
        f'<w:insideH w:val="single" w:sz="4" w:color="BFCBD7"/>'
        f'<w:insideV w:val="single" w:sz="4" w:color="BFCBD7"/>'
        f'</w:tblBorders>')
    tblPr.append(borders)
    return table

def section_heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    run.bold = True; run.font.size = Pt(13); run.font.color.rgb = DARK_BLUE; run.font.name = 'Calibri'
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    pPr = p2._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:color="0B2A4A" w:space="4"/></w:pBdr>')
    pPr.append(pBdr)

def sub_heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.bold = True; run.font.size = Pt(11); run.font.color.rgb = MED_BLUE; run.font.name = 'Calibri'

def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.size = Pt(10); run.font.name = 'Calibri'

def bullet(doc, text, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0 + indent)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(f"• {text}")
    run.font.size = Pt(10); run.font.name = 'Calibri'

def urgent_box(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(f"⚠ URGENT / PENDING: {text}")
    run.bold = True; run.font.size = Pt(10); run.font.color.rgb = RED_URGENT; run.font.name = 'Calibri'

# ═══════════════════════════════════════════════════════════════════════════════
# DOCUMENT HEADER
# ═══════════════════════════════════════════════════════════════════════════════

p_bar = doc.add_paragraph()
pPr = p_bar._p.get_or_add_pPr()
pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="36" w:color="0B2A4A" w:space="0"/></w:pBdr>')
pPr.append(pBdr)
doc.add_paragraph("")

p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p_title.add_run("WORK HANDOVER MEMO — TEMPORARY LEAVE")
run.bold = True; run.font.size = Pt(18); run.font.color.rgb = DARK_BLUE; run.font.name = 'Calibri'

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_after = Pt(12)
run = p_sub.add_run("Commercial / Contracts Management — International EPC")
run.font.size = Pt(11); run.font.color.rgb = ACCENT_GREY; run.font.name = 'Calibri'

handover_info = [
    ["Prepared By", "[Name]", "Date", today_str],
    ["Leave Period", "[DD-Mon-YYYY]  to  [DD-Mon-YYYY]", "Contact (Urgent)", "[Phone / WeChat / Email]"],
    ["Backup Contact", "[Colleague Name / Phone]", "Handover Receiver", "[Name of person receiving this handover]"],
]
add_table(doc, handover_info[0], handover_info[1:], col_widths=[3.5, 5.0, 3.5, 5.0])

doc.add_paragraph("")

# ═══════════════════════════════════════════════════════════════════════════════
# 1. SYSTEM & INFRASTRUCTURE
# ═══════════════════════════════════════════════════════════════════════════════

section_heading(doc, "1. SYSTEM & INFRASTRUCTURE (Daily Operations)")

sub_heading(doc, "1.1  File & Directory Rules")
body(doc, "Three-location segregation — new software, scripts, and templates each go to their designated location. Violating this causes clutter and data loss.")
rules_table = [
    ["Installed Software", r"D:\Program Files\<AppName>"],
    ["Wison Project Tools & Scripts", r"D:\Wison\_tools  (git-managed)"],
    ["System / Dev Config Scripts", r"D:\Documents\My Projects"],
    ["Wison Templates (daily use)", r"D:\Wison\Project_Info\Wison Template"],
    ["Non-Wison Templates", r"D:\Wison\Project_Info\Wison Template\Non-Wison Template"],
    ["Temp Files (project/delivery areas)", r"D:\Documents root only; NEVER D:\Temp"],
]
add_table(doc, ["Category", "Location"], rules_table, col_widths=[6.0, 10.5])

sub_heading(doc, "1.2  Deletion Policy (CRITICAL)")
body(doc, "ALL deletions go to Windows Recycle Bin — NEVER permanent delete. PowerShell: use Microsoft.VisualBasic.FileIO.FileSystem.DeleteFile/DeleteDirectory with 'SendToRecycleBin'. Never use Remove-Item -Force.")
body(doc, "Recycle Bin auto-purges after 30 days (Storage Sense policy at HKCU StoragePolicy: 08=1, 256=30). Two-stage: Temp >30d → Recycle Bin; Recycle Bin >30d → permanent. Max ~60-day recovery window.")

sub_heading(doc, "1.3  Backup System")
body(doc, r"Location: D:\Program Files\Backup")
bullet(doc, "backup-all.ps1 — Full backup script (run before major changes)")
bullet(doc, "RESTORE.ps1 — Restoration script (read before using)")
bullet(doc, r"Weasel/RIME input method — backs up from live RimeUserDir automatically")
bullet(doc, r"Anki — app at D:\Program Files\Anki (junction), data at D:\Documents\Anki\Anki2 (junction off %APPDATA%); data syncs via AnkiWeb cloud — no local backup needed")

sub_heading(doc, "1.4  Scheduled Tasks")
body(doc, "Windows Task Scheduler has an automated cleanup task:")
bullet(doc, "Task: Clean-UserTemp-DDocuments — runs daily, cleans D:\\Documents\\Temp, files go to Recycle Bin (30-day retention)")
bullet(doc, r"Script: D:\Documents\My Projects\scripts\system\clean-temp.ps1")
bullet(doc, r"Storage Sense policy: daily check, Temp files >30d → Recycle Bin, never bypasses Recycle Bin")

sub_heading(doc, "1.5  Claude Code CLI")
body(doc, "npm global install, version-locked at 2.1.207 (no auto-update). Codebase/global instructions at C:\\Users\\Admin\\.claude\\CLAUDE.md. Memory files at C:\\Users\\Admin\\.claude\\projects\\D--Documents-My-Projects\\memory\\. Session resumption via claude --resume <UUID>.")
bullet(doc, "CLAUDE.md — global work conventions (file routing, deletion policy, temp file rule, etc.)")
bullet(doc, "MEMORY.md — index of all saved facts; each .md file = one fact with frontmatter")
bullet(doc, "Session index at memory/session-index.md — 10 preserved session UUID→topic mappings")

# ═══════════════════════════════════════════════════════════════════════════════
# 2. WISON PROJECT KNOWLEDGE BASE
# ═══════════════════════════════════════════════════════════════════════════════

section_heading(doc, "2. WISON PROJECT — KNOWLEDGE BASE & ACTIVE WORK")

sub_heading(doc, r"2.1  D:\Wison\ Directory Structure")
body(doc, "Six subdirectories with distinct purposes. Do NOT mix them.")
wison_dirs = [
    [r"Project_Info", "Project meta-information, templates, reference documents"],
    [r"_tools", "All tools, scripts, guides (git-managed). Payment app generators, OCR scripts, pricing DB tools"],
    [r"原始档案 (Original Archives)", "Raw contract documents — read-only reference"],
    [r"AI衍生数据 (AI-Derived)", "AI-processed outputs, extracted data"],
    [r"知识库 (Knowledge Base)", "Structured knowledge derived from project documents"],
    [r"交付物 (Deliverables)", "Final deliverables (.docx, .xlsx, etc.)"],
]
add_table(doc, ["Directory", "Purpose"], wison_dirs, col_widths=[6.0, 10.5])

sub_heading(doc, "2.2  Active Wison Subcontracts & Commercial Work")
body(doc, "The following are active/known workstreams that the handover receiver needs to be aware of:")

contracts_table = [
    ["12.1 CCECC-MEI Pkg I", "Amendment 01 (Amdt 01) — finalised after 15 drafting iterations. Amdt 01 EN alignment done. LOA attachments renamed to Pkg system (body text untouched). Descope financial logic applied."],
    ["12.1 Sub-packages", "PK-3 = Pkg I-3 (北方国际 + 二十三冶). TCC = 三化建 (NOT 中交天航). MEI subcontractor→package mapping documented."],
    ["SGP / RSHT", "SGP (not RSGP) + RSHT-2 naming convention with spaces. Full name first use, then abbreviation."],
    ["Payment Applications", "Standardised template created (see Section 3 below)."],
    ["BOQ Pricing Database", r"135 CSVs across 7 subdirectories under D:\Wison\_tools\price-database\. All prices/rates/totals are excl. VAT by default. Covers Civil + MEI + fire protection + prefabrication."],
]
add_table(doc, ["Subcontract / Area", "Status & Key Notes"], contracts_table, col_widths=[4.5, 12.0])

sub_heading(doc, "2.3  Contract Review & Drafting Conventions")
body(doc, "Established conventions that MUST be followed when reviewing or drafting any Wison contract document:")

conventions = [
    ["Amendment Drafting", "Tri-party structure, clause numbering system, CN/EN bilingual layout, Descope financial logic, signature pages, abbreviation format (remove brackets). Ref: amdt-drafting-conventions memory."],
    ["Contract Review", "Auto-check on every review: spelling, terminology consistency (FIDIC 1999/2017), substantive clause issues. Always create backup before modifying. Ref: contract-review-rules memory."],
    ["Letter Writing", "CN/EN bilingual demand/notice letters: concise wording, logical structure, clean style. Based on 12.1 Delay Notice user revisions. Ref: wison-letter-writing-style memory."],
    ["Naming Protocol", "12.1 CCECC-MEI I is directory template. Filenames NOT to be changed (especially ATT/Exhibit sub-files). LOA/Amdt per actual circumstances. Ref: wison-naming-protocol memory."],
    ["FIDIC Terminology", "Contract ≠ Contract Agreement. Employer/Contractor roles. Rainbow Suite. Document priority order. Ref: fidic-terminology memory."],
]
add_table(doc, ["Area", "Rule"], conventions, col_widths=[4.0, 12.5])

# ═══════════════════════════════════════════════════════════════════════════════
# 3. RECENTLY CREATED TEMPLATES & TOOLS
# ═══════════════════════════════════════════════════════════════════════════════

section_heading(doc, "3. RECENTLY CREATED — TEMPLATES & TOOLS (July 2026)")

sub_heading(doc, "3.1  Subcontractor Application for Payment Template")
body(doc, "Created 23-Jul-2026. Universal payment application for deliverable-based / professional service subcontracts (NOT construction BOQ-based). Two formats available:")

pay_template = [
    ["Excel (.xlsx) ★ Recommended", r"D:\Wison\Project_Info\Wison Template\Non-Wison Template\Subcontractor_Application_for_Payment_Template.xlsx", "10 sheets: Cover, Project Info, Commercial Summary (formulas), Deliverables (dropdown validation), Milestones (SUM weight auto-calc), Supporting Docs (checkboxes + attachment register), Subcontractor Certification (7 warranty statements), Contractor Review (10 disciplines), Signatures (Subcontractor blue / EPC gold), Guidance Notes"],
    ["Word (.docx)", r"D:\Wison\Project_Info\Wison Template\Non-Wison Template\Subcontractor_Application_for_Payment_Template.docx", "8 sections: same content, print-optimised, no formulas"],
    ["Generator Scripts", r"D:\Wison\_tools\generate_payment_application.py  &  generate_payment_application_xlsx.py", "Run these to regenerate templates after modifications. Output paths point to Non-Wison Template directory above."],
]
add_table(doc, ["Item", "Location", "Notes"], pay_template, col_widths=[3.5, 6.5, 6.5])

sub_heading(doc, "3.2  ADNOC Sulfur Guidebook — QA/QC Alignment Script")
body(doc, "Python script for automatic engineering QA/QC alignment of the ADNOC sulfur guidebook HTML. Started v6, upgrade to v7 in progress (revision-based optimisation).")
bullet(doc, "Location: (in user's current working files — check recent Claude Code session)")
bullet(doc, "Functions: contractual reference rectification (Clause 3(b) → 3.2(b)), Clause 5.1 flattening (remove sub-levels), legal source attribution injection in parentheses")

sub_heading(doc, "3.3  OCR & Document Processing Pipeline")
body(doc, "Two-method pipeline for converting documents to Markdown:")
bullet(doc, "Method 2: MarkItDown MCP (preferred for simple docs)")
bullet(doc, "Method 3: Custom Python pipeline (pymupdf/fitz render → Tesseract OCR)")
bullet(doc, "No poppler dependency — pymupdf (fitz) for rendering")
bullet(doc, "Excel/CSV protocol: .xlsx/.xls → one CSV per sheet, same base name, metadata header")
bullet(doc, "Tools location: D:\\Wison\\_tools\\ocr-tools\\pdf-toolbox\\ (see README.md)")

sub_heading(doc, "3.4  PFD Visual Style (Stacked HTML)")
body(doc, "Final PFD colour scheme: Twilight dark industrial, jacket section panels, 4 material layer leader lines, light print version rules. Ref: pfd-visual-style memory.")

# ═══════════════════════════════════════════════════════════════════════════════
# 4. PENDING / UNFINISHED ITEMS
# ═══════════════════════════════════════════════════════════════════════════════

section_heading(doc, "4. PENDING / UNFINISHED ITEMS — NEED ATTENTION")

urgent_box(doc, "Wison-Fabricator Agreement Simplification — DO NOT EXECUTE YET. 3-part simplification plan drafted (what to delete, what to keep). Awaiting user go-ahead. Ref: wison-agreement-simplification memory.")
urgent_box(doc, "Windows Security Center Dead Icon (SEC Health) — NOT FIXED. Root cause: corrupted zh-CN PRI file, TrustedInstaller locked. Workaround path not yet found. Ref: sec-health-dead-icon memory.")
urgent_box(doc, "ADNOC Sulfur Guidebook v7 Upgrade — In progress. Revision-based optimisation of the QA/QC alignment script. v6 tested, v7 pending finalisation.")
urgent_box(doc, "drive263Uploader — Installed to D:\\Program Files\\drive263Uploader\\. NPAPI plugin, NOT an executable. Edge does NOT load it. Website (drive.263.net) works without it. Plugin effectively unused — may be removed.")

# ═══════════════════════════════════════════════════════════════════════════════
# 5. RECURRING WORKFLOWS & CHECKS
# ═══════════════════════════════════════════════════════════════════════════════

section_heading(doc, "5. RECURRING WORKFLOWS & PERIODIC CHECKS")

recurring = [
    ["Daily", "TEMP cleanup", "Automatic via Task Scheduler (Clean-UserTemp-DDocuments). Files >30d → Recycle Bin. No manual action needed."],
    ["Daily", "Recycle Bin check", "Storage Sense auto-purges items >30 days. Critical files in Recycle Bin should be restored before 30-day limit."],
    ["Weekly", "Backup", r"Run D:\Program Files\Backup\backup-all.ps1 (or verify last run date)."],
    ["Per Contract Review", "Spelling / terminology consistency check", "Auto-check on every contract document review — FIDIC terminology, clause numbering, CN/EN alignment."],
    ["Per Payment Application", "Template usage", "Use Subcontractor_Application_for_Payment_Template.xlsx in Non-Wison Template. Commercial Summary has auto-calc formulas — do NOT overwrite shaded cells."],
    ["Per Amendment", "Drafting conventions", "Follow amdt-drafting-conventions: tri-party structure, CN/EN bilingual, Descope logic, signature pages, abbreviation format."],
    ["As needed", "Claude Code session resume", "claude --resume <UUID> from session-index.md memory file. 10 sessions preserved."],
]
add_table(doc, ["Frequency", "Task", "Notes"], recurring, col_widths=[3.0, 4.0, 9.5])

# ═══════════════════════════════════════════════════════════════════════════════
# 6. KEY FILE PATHS QUICK REFERENCE
# ═══════════════════════════════════════════════════════════════════════════════

section_heading(doc, "6. KEY FILE PATHS — QUICK REFERENCE")

paths_ref = [
    ["Global Work Conventions", r"C:\Users\Admin\.claude\CLAUDE.md"],
    ["Memory / Knowledge Index", r"C:\Users\Admin\.claude\projects\D--Documents-My-Projects\memory\MEMORY.md"],
    ["Claude Code Sessions", r"C:\Users\Admin\.claude\projects\D--Documents-My-Projects\memory\session-index.md"],
    ["Backup Scripts", r"D:\Program Files\Backup"],
    ["TEMP Cleanup Script", r"D:\Documents\My Projects\scripts\system\clean-temp.ps1"],
    ["Wison Tools (git)", r"D:\Wison\_tools"],
    ["Wison Templates", r"D:\Wison\Project_Info\Wison Template"],
    ["Payment App Template (Excel)", r"D:\Wison\Project_Info\Wison Template\Non-Wison Template\Subcontractor_Application_for_Payment_Template.xlsx"],
    ["Payment App Template (Word)", r"D:\Wison\Project_Info\Wison Template\Non-Wison Template\Subcontractor_Application_for_Payment_Template.docx"],
    ["Payment App Generators", r"D:\Wison\_tools\generate_payment_application_xlsx.py  (and .py)"],
    ["BOQ Price Database", r"D:\Wison\_tools\price-database\  (135 CSVs, 7 sub-dirs)"],
    ["OCR Tools", r"D:\Wison\_tools\ocr-tools\pdf-toolbox\README.md"],
    ["Contract Review Rules", r"C:\Users\Admin\.claude\projects\D--Documents-My-Projects\memory\contract-review-rules.md"],
    ["Amendment Conventions", r"C:\Users\Admin\.claude\projects\D--Documents-My-Projects\memory\amdt-drafting-conventions.md"],
    ["MEI Subcontractor Mapping", r"C:\Users\Admin\.claude\projects\D--Documents-My-Projects\memory\mei-subcontractor-mapping.md"],
]
add_table(doc, ["What", "Path"], paths_ref, col_widths=[5.5, 11.0])

# ═══════════════════════════════════════════════════════════════════════════════
# 7. HANDOVER NOTES (Free-text)
# ═══════════════════════════════════════════════════════════════════════════════

section_heading(doc, "7. HANDOVER NOTES & CONTEXT")

body(doc, "The following notes provide context that may not be obvious from the file system or tools alone:")

body(doc, "7.1  How Claude Code / AI Assistant Works in This Environment")
bullet(doc, "The AI assistant (Claude Code) has persistent memory of all conventions, rules, and past work via the memory system at C:\\Users\\Admin\\.claude\\projects\\D--Documents-My-Projects\\memory\\. It reads CLAUDE.md (global conventions) and MEMORY.md (fact index) at the start of every session.")
bullet(doc, "When asking the AI to do work, it will auto-recall relevant memories. No need to re-explain conventions — they are permanently stored.")
bullet(doc, "The AI CAN generate documents, scripts, templates, and perform contract review. It CANNOT access the internet unless explicitly directed, and CANNOT log into online services without user credentials.")
bullet(doc, "Session resumption: claude --resume <UUID> to pick up where a previous session left off. UUIDs are logged in memory/session-index.md.")

body(doc, "7.2  Critical Dos and Don'ts")
bullet(doc, "DO always use absolute paths in scripts (scripts run from anywhere).")
bullet(doc, "DO send deletions to Recycle Bin (never permanent).")
bullet(doc, "DO backup before major changes (run backup-all.ps1).")
bullet(doc, "DO follow amendment drafting conventions for any contract variation (tri-party, CN/EN, Descope logic).")
bullet(doc, "DON'T mix file categories: Wison project work → Wison\\_tools, system config → My Projects, software → Program Files.")
bullet(doc, "DON'T create temp files in D:\\Temp or project directories — use D:\\Documents\\ root.")
bullet(doc, "DON'T execute the Wison-Fabricator agreement simplification without explicit user instruction.")
bullet(doc, "DON'T modify Wison original contract archive files (原始档案) — read-only reference.")

body(doc, "7.3  If Something Breaks")
bullet(doc, "Check Windows Recycle Bin first — most deleted files are recoverable within 30 days.")
bullet(doc, "Restore from backup: D:\\Program Files\\Backup\\RESTORE.ps1.")
bullet(doc, "Git history: D:\\Wison is a git repo — git log / git diff for tracking changes to tools and knowledge base.")
bullet(doc, "Claude Code memory: all conventions and rules are at C:\\Users\\Admin\\.claude\\ — this is the AI's 'brain' for this machine.")
bullet(doc, "Weasel/RIME input method backup: live backup from RimeUserDir; restore procedure in backup-system.md memory.")

# ═══════════════════════════════════════════════════════════════════════════════
# SIGN-OFF
# ═══════════════════════════════════════════════════════════════════════════════

doc.add_paragraph("")
section_heading(doc, "ACKNOWLEDGMENT OF RECEIPT")

body(doc, "I acknowledge receipt of this Handover Memo and confirm that I have been briefed on all items listed above. I understand the urgency of pending items marked in red and the criticality of the file management and deletion policies.")

sig_headers = ["", "Name", "Signature", "Date"]
sig_rows = [
    ["Handover By\n(Outgoing)", "[Name]", "", "[DD-Mon-YYYY]"],
    ["Received By\n(Incoming / Covering)", "[Name]", "", "[DD-Mon-YYYY]"],
]
add_table(doc, sig_headers, sig_rows, col_widths=[5.0, 4.0, 3.5, 4.0])

doc.add_paragraph("")
p_end = doc.add_paragraph()
p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p_end.add_run("— END OF HANDOVER MEMO —")
run.bold = True; run.font.size = Pt(9); run.font.color.rgb = ACCENT_GREY; run.font.name = 'Calibri'

# ── Save ─────────────────────────────────────────────────────────────────────
output = r"D:\Wison\Project_Info\Wison Template\Non-Wison Template\Work_Handover_Memo_Template.docx"
doc.save(output)
print(f"[DONE] Handover memo saved to: {output}")
import os; os.startfile(output)
