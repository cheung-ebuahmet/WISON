"""
Generate a minimal one-page Service Instruction for additional night-shift
medical personnel under the existing Medical Service Agreement.

This is NOT a Purchase Order or contract amendment. It is a simple written
instruction so the Supplier has a basis for monthly invoicing.
All rates, payment terms, and contractual provisions remain per the
existing Agreement.
"""
import os, sys
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', errors='replace')

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

OUT = r"D:\Wison\Subcon_Payments\08 Med. Svc. Agr\Contract\Medical_Service_Additional_Personnel_Request.docx"

doc = Document()

# ---- Page Setup ----
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(4)

# ---- Helpers ----
BLUE = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x80, 0x80, 0x80)

def p(text, bold=False, size=11, align=None, color=None, sb=0, sa=4):
    """Add a paragraph."""
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(sb)
    par.paragraph_format.space_after = Pt(sa)
    if align is not None:
        par.alignment = align
    run = par.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return par

def hr():
    """Thin grey horizontal rule."""
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(2)
    pPr = par._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="1" w:color="999999"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)

def tbl_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else \
        parse_xml(f'<w:tblPr {nsdecls("w")}></w:tblPr>')
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="333333"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="333333"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="333333"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="333333"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="333333"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="333333"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def cell_shade(cell, color):
    elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(elm)

def cell_pad(cell, t=50, b=50, l=100, r=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}></w:tcMar>')
    for m, v in [('top',t), ('bottom',b), ('left',l), ('right',r)]:
        nd = parse_xml(f'<w:{m} {nsdecls("w")} w:w="{v}" w:type="dxa"/>')
        tcMar.append(nd)
    tcPr.append(tcMar)

def cell_write(cell, text, bold=False, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell_pad(cell)
    par = cell.paragraphs[0]
    par.alignment = align
    run = par.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.bold = bold

# ============================================================
# LETTERHEAD
# ============================================================
p("WISON ENERGY ENGINEERING (HONG KONG) LIMITED – ABU DHABI",
  bold=True, size=13, color=BLUE, sa=1)
p("Saman Tower, Floor 13, Office 1309, Hamdan Street, Abu Dhabi, UAE",
  size=8.5, color=GREY, sa=0)
p("Tel: +971 2 681 5588  |  Project: Sulphur Granulation Plant at RSHT-2",
  size=8.5, color=GREY, sa=8)
hr()

# ============================================================
# TO / DATE / REFERENCE
# ============================================================
p("", sa=2)
p("To:     RESPONSE PLUS MEDICAL SERVICES L.L.C.", bold=False, size=10.5, sa=1)
p("         6th Floor, Emirates Real Estate Corp Building,", size=9.5, color=GREY, sa=0)
p("         Al Falah Street, Al Danah, Abu Dhabi. P.O.Box 130336", size=9.5, color=GREY, sa=6)

p("Date:   ____________________", size=10.5, sa=4)
p("Ref.:   Medical Service Agreement — Night Shift Additional Personnel", size=10.5, sa=10)
hr()

# ============================================================
# TITLE
# ============================================================
p("REQUEST FOR ADDITIONAL MEDICAL PERSONNEL",
  bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, color=BLUE, sb=10, sa=12)
hr()

# ============================================================
# BODY
# ============================================================
p("Pursuant to the Medical Service Agreement between the parties, "
  "Wison hereby requests the deployment of the following additional "
  "personnel for night shift services:",
  size=11, sb=10, sa=10)

# --- Positions Table ---
t = doc.add_table(rows=3, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
tbl_borders(t)

# Header row
hdr = t.rows[0]
for i, (txt, w) in enumerate(zip(
    ["Position", "Qty", "Shift"],
    [9.5, 2.0, 4.5])):
    c = hdr.cells[i]
    cell_shade(c, '1F4E79')
    cell_pad(c, t=50, b=50, l=80, r=80)
    c.width = Cm(w)
    pa = c.paragraphs[0]
    pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = pa.add_run(txt)
    r.font.name = 'Calibri'; r.font.size = Pt(10); r.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

# Data rows
cell_write(t.rows[1].cells[0], "Registered Male Nurse (DOH-Approved, BLS/ACLS Certified)")
cell_write(t.rows[1].cells[1], "1", align=WD_ALIGN_PARAGRAPH.CENTER)
cell_write(t.rows[1].cells[2], "Night Shift   ·   10 hrs/day   ·   6 days/week",
           align=WD_ALIGN_PARAGRAPH.CENTER)

cell_write(t.rows[2].cells[0], "First Aid Trained Driver (Valid UAE Driving License)")
cell_write(t.rows[2].cells[1], "1", align=WD_ALIGN_PARAGRAPH.CENTER)
cell_write(t.rows[2].cells[2], "Night Shift   ·   10 hrs/day   ·   6 days/week",
           align=WD_ALIGN_PARAGRAPH.CENTER)

for row in [t.rows[1], t.rows[2]]:
    for c in row.cells:
        cell_shade(c, 'F2F5FA')

p("", sa=12)

# ============================================================
# COMMERCIAL — One paragraph
# ============================================================
hr()
p("The applicable unit rates, payment terms, and all other contractual "
  "provisions shall remain in accordance with the existing Medical Service "
  "Agreement. Payment shall be based on actual attendance records, verified "
  "and approved by the Contractor on a monthly basis.",
  size=10.5, sb=10, sa=10)
hr()

# ============================================================
# SIGN-OFF — Issuer only, no acknowledgement
# ============================================================
p("", sa=6)
p("Issued by:", bold=True, size=10.5, sa=6)
p("_________________________________", size=10, sa=1)
p("Name", size=9, color=GREY, sa=8)
p("_________________________________", size=10, sa=1)
p("Title", size=9, color=GREY, sa=8)
p("_________________________________", size=10, sa=1)
p("Date", size=9, color=GREY, sa=2)

# ============================================================
# SAVE
# ============================================================
doc.save(OUT)
print(f"Saved: {OUT}")
