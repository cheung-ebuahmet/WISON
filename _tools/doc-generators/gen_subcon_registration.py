# -*- coding: utf-8 -*-
"""
Subcontractor Registration (REVISED) — generator
=================================================
Single sheet, modules stacked vertically (A, B, C ...).
Global top: title → legend → header. Then each module = title band + rows.

Row types (in order):
  'fill'     → white/zebra row, Response empty (dropdown/date/number/text)
  'upload'   → yellow row, No. = A1/A2/... (sub-ref → folder)
  'subhead'  → light-blue band (section divider)
  'note'     → grey italic instruction line

Language: ENGLISH ONLY.
"""
import glob
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

DARK    = '1F4E79'
MID     = '2E75B6'
SUBHEAD = 'BDD7EE'
LIGHT   = 'DDEBF7'
UP      = 'FFF2CC'
REQ     = 'C00000'
GREY    = '595959'
GRAYTX  = '7F7F7F'

thin   = Side(style='thin', color='BFBFBF')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def resolve_dir():
    root = r'D:/Wison/Project_Info/Wison Template'
    parent = None
    for d in glob.glob(root + '/*'):
        b = os.path.basename(d)
        if 'Pre-Qualification' in b and 'ADNOC' in b:
            parent = d
            break
    if parent is None:
        raise SystemExit('parent dir not found')
    out = os.path.join(parent, 'ADNOC WISON SGP Pre-Qualification Requirements (REVISED)')
    os.makedirs(out, exist_ok=True)
    return out

MODULES = [
    {
        'id': 'A', 'title': 'Company Information', 'folder': 'A - Company Information',
        'rows': [
            {'type':'upload','item':'Business Registration Certificate','required':True},
            {'type':'fill','item':'Company name','required':False,'ftype':'text'},
            {'type':'fill','item':'Contact address','required':False,'ftype':'text_large'},
            {'type':'fill','item':'Business Registration Certificate No.','required':False,'ftype':'text'},
            {'type':'fill','item':'Nature of business','required':False,'ftype':'dropdown',
             'options':['Corporation','Sole Proprietor','LLC','Partnership','General or Limited','Joint Venture']},
            {'type':'fill','item':'Established date','required':False,'ftype':'date'},
            {'type':'fill','item':'President / Owner / Partner name','required':True,'ftype':'text'},
            {'type':'fill','item':'Contact email','required':True,'ftype':'text','key':'contact_email'},
            {'type':'upload','item':'Type of Work / Activities','required':True,
             'note':'Qualification & Licenses Certificate. In UAE usually the Chamber of Commerce & '
                    'Industry certificate; if not available, re-upload the Business Registration '
                    'Certificate instead.'},
            {'type':'fill','item':'Type of Work / Activities','required':True,'ftype':'dropdown',
             'share_title':True, 'note':'Select work type',
             'options':['Temporary Facilities','Site Preparation','Pile','Civil','Building','Road',
                        'Underground Pipe','Equipment Pipe','Steel Structure','Tank','Electric & Instrument',
                        'Heavy Lifting','Paint & Insulation','Scaffold','Fire Fighting','Fire Proofing',
                        'Chemical Cleaning','Nondestructive Examination','Other']},
            {'type':'upload','item':'Company profile','required':True,
             'note':'Company brochure / profile (PDF, PPT, etc.) as available'},
        ],
    },
    {
        'id': 'B', 'title': 'Legal Compliance', 'folder': 'B - Legal Compliance',
        'rows': [
            {'type':'note','text':'Use the latest template downloaded from the system (do not modify). '
                                 'For an electronic seal/signature, attach the platform seal process or '
                                 'verification link (cut-and-paste is invalid). Also attach proof the '
                                 'signatory is an employee (Proof of Employment / Company-Verified '
                                 'LinkedIn / Residence ID, etc.).'},
            {'type':'upload','item':'Due Diligence Questionnaire (stamped)','required':True,'template':True},
            {'type':'upload','item':'Undertaking of Honest Conduct (stamped)','required':True,'template':True},
        ],
    },
    {
        'id': 'C', 'title': 'Financial Information', 'folder': 'C - Financial Information',
        'rows': [
            {'type':'subhead','text':'Part 1 — Financial Indicators'},
            {'type':'upload','item':'Certified financial statement of the last three years','required':True},
            {'type':'fill','item':'Currency','required':False,'ftype':'dropdown','default':'AED',
             'options':['AED','USD','CNY','EUR','GBP','SAR','OMR','KWD','BHD','QAR','AZN','BRL','Other'],
             'note':'Default AED. You may select a different currency if required.'},
            {'type':'fill','item':'Total Assets (ten thousand)','required':False,'ftype':'number','note':'ten thousand'},
            {'type':'fill','item':'Net Asset Amount (ten thousand)','required':False,'ftype':'number','note':'ten thousand'},
            {'type':'fill','item':'Revenue (ten thousand)','required':False,'ftype':'number','note':'ten thousand'},
            {'type':'fill','item':'Net Profit (ten thousand)','required':False,'ftype':'number','note':'ten thousand'},
            {'type':'fill','item':'Debt to Asset Ratio','required':False,'ftype':'number','note':'percentage %'},
            {'type':'fill','item':'Net Profit Margin','required':False,'ftype':'number','note':'percentage %'},

            {'type':'subhead','text':'Part 2 — Bank Account'},
            {'type':'upload','item':'Supplier Receipt Account Confirmation Letter','required':True,
             'note':'With company stamp & financial department stamp','template':True},
            {'type':'upload','item':'Bank account proof (bank issued)','required':True,
             'note':'With company stamp & financial department stamp','template':'reference'},
            {'type':'fill','item':'Account type','required':False,'ftype':'dropdown',
             'options':['T/T (Telegraphic Transfer)','Electronic Acceptance','T/T and Electronic Acceptance'],
             'note':'default T/T. You may select a different type if required.'},
            {'type':'fill','item':'Bank account name','required':False,'ftype':'text','note':'usually same as company name'},
            {'type':'fill','item':'Bank of deposit','required':False,'ftype':'text'},
            {'type':'fill','item':'Bank address','required':False,'ftype':'text'},
            {'type':'fill','item':'Account number','required':False,'ftype':'text'},
            {'type':'fill','item':'Swift Code','required':False,'ftype':'text'},
            {'type':'fill','item':'Payee address','required':False,'ftype':'text'},
            {'type':'fill','item':'Payee city','required':False,'ftype':'text'},
        ],
    },
    {
        'id': 'D', 'title': 'Contact of Supplier', 'folder': 'D - Contact of Supplier',
        'rows': [
            {'type':'note','text':'Enter the supplier contact person (company personnel). Set one contact '
                                 'as the platform login account — it is auto-generated after approval.'},
            {'type':'upload','item':'Authorization Letter — Supplier Portal Account','required':True,
             'note':'With company stamp and corresponding ID scan (you can watermark for security, etc.)','template':True},
            {'type':'fill','item':'Position','required':False,'ftype':'dropdown',
             'options':['Technical in charge','Marketing in charge','Others']},
            {'type':'fill','item':'Full name','required':False,'ftype':'text'},
            {'type':'fill','item':'E-mail','required':False,'ftype':'text','ref':'contact_email'},
            {'type':'upload','item':'Authorization Letter — Business Leader','required':True,
             'note':'With company stamp and corresponding ID scan (you can watermark for security, etc.)','template':True},
        ],
    },
    {
        'id': 'E', 'title': 'Qualification Information', 'folder': 'E - Qualification Information',
        'rows': [
            {'type':'fill','item':'Quality Standard as applied by you','required':True,'ftype':'dropdown',
             'options':['ASME','SCA','Other']},
            {'type':'upload','item':'Attachment (QA Certificate)','required':True,
             'note':'Upload the QA certificate matching the selected Quality Standard'},
            {'type':'fill','item':'Certificate number','required':False,'ftype':'text'},
            {'type':'fill','item':'Certificate date','required':False,'ftype':'date'},
            {'type':'fill','item':'Quality System subject to independent 3rd party assessment','required':True,
             'ftype':'dropdown','options':['Yes','No']},
            {'type':'upload','item':'Approval Certificate / Quality Manual','required':True,
             'note':'If Yes: upload 3rd Party Approval Certificate & Scope of Approval. '
                    'If No: upload Quality Manual copy. (Either way — mandatory)'},
        ],
    },
    {
        'id': 'F', 'title': 'ICV Certificate', 'folder': 'F - ICV Certificate',
        'rows': [
            {'type':'fill','item':'ICV certificate','required':True,'ftype':'dropdown',
             'options':['Yes','No']},
            {'type':'note','emph':True,
             'text':'If you select YES → the attachment below MUST be uploaded. '
                    'If NO → skip this module.'},
            {'type':'upload','item':'ICV Certificate (attachment)','required':True,
             'note':'Upload ICV Certificate scan'},
            {'type':'fill','item':'Certificate number','required':False,'ftype':'text'},
            {'type':'fill','item':'Issue date','required':False,'ftype':'date'},
            {'type':'fill','item':'Valid date','required':False,'ftype':'date'},
            {'type':'fill','item':'ICV Score','required':False,'ftype':'number'},
        ],
    },
    {
        'id': 'G', 'title': 'Construction Capacity', 'folder': 'G - Construction Capacity',
        'rows': [
            {'type':'subhead','text':'Ability to accept contract amount'},
            {'type':'fill','item':'Min. ($)','required':False,'ftype':'number','default':0,'note':'USD $'},
            {'type':'fill','item':'Max. ($)','required':False,'ftype':'number','default':0,'note':'USD $'},
            {'type':'subhead','text':'Workforce'},
            {'type':'fill','item':'Skilled Workers Total Number of People','required':False,'ftype':'number','default':0},
            {'type':'fill','item':'Semi-Skilled Workers','required':False,'ftype':'number','default':0},
            {'type':'fill','item':'Helpers','required':False,'ftype':'number','default':0},
            {'type':'fill','item':'Supervision (Foremen/General Foremen)','required':False,'ftype':'number','default':0},
            {'type':'fill','item':'Supervision (Supervisors)','required':False,'ftype':'number','default':0},
            {'type':'fill','item':'Professional (Safety/Scheduling/QC)','required':False,'ftype':'number','default':0},
            {'type':'fill','item':'Administration/Management','required':False,'ftype':'number','default':0},
            {'type':'fill','item':'Total Workforce','required':False,'ftype':'number','default':0,'auto_sum':True},
        ],
    },
    {
        'id': 'H', 'title': 'Construction Equipment and Tools list', 'folder': 'H - Construction Equipment and Tools',
        'rows': [
            {'type':'fill','item':'Mechanical equipment classification','required':False,'ftype':'dropdown',
             'options':['Hoisting Machinery','Pile Foundation Machinery','Construction Machine',
                        'Welding Equipment','Test Instrument','Others']},
            {'type':'fill','item':'Mechanical name','required':False,'ftype':'text'},
            {'type':'fill','item':'Rule model','required':False,'ftype':'text'},
            {'type':'fill','item':'Quantity','required':False,'ftype':'number'},
        ],
    },
    {
        'id': 'I', 'title': 'Experience', 'folder': 'I - Experience',
        'rows': [
            {'type':'fill','item':'Name of project','required':True,'ftype':'text'},
            {'type':'fill','item':'Client/Owner','required':True,'ftype':'text'},
            {'type':'fill','item':'General Contractor','required':True,'ftype':'text'},
            {'type':'fill','item':'Contract Value ($)','required':True,'ftype':'number','note':'USD $'},
            {'type':'fill','item':'Completion Date','required':True,'ftype':'date'},
            {'type':'fill','item':'Description of Work Being Performed','required':True,'ftype':'text_large'},
        ],
    },
]

def build():
    out_dir = resolve_dir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Registration'
    for col, w in {'A':7,'B':42,'C':34,'D':48}.items():
        ws.column_dimensions[col].width = w

    # ---- 隐藏 Lists sheet：承载所有下拉选项（避免 Excel 内联列表 255 字符上限）----
    lists_ws = wb.create_sheet('Lists')
    lists_ws.sheet_state = 'hidden'
    opt_name = {}   # tuple(options) -> named range name
    opt_col = {}    # tuple(options) -> column letter
    col_idx = 1
    for m in MODULES:
        for r in m['rows']:
            if r.get('type') == 'fill' and r.get('ftype') == 'dropdown':
                opts = tuple(r.get('options', []))
                if opts and opts not in opt_name:
                    cl = get_column_letter(col_idx)
                    nm = f'List{col_idx}'
                    opt_name[opts] = nm
                    opt_col[opts] = cl
                    for i, val in enumerate(opts, start=1):
                        lists_ws.cell(row=i, column=col_idx, value=val)
                    # 命名区域指向 Lists 列
                    wb.defined_names.add(
                        DefinedName(name=nm, attr_text=f"Lists!${cl}$1:${cl}${len(opts)}"))
                    col_idx += 1

    row = 1
    # top title
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    c = ws.cell(row=row, column=1, value='SUBCONTRACTOR REGISTRATION — PRE-QUALIFICATION')
    c.font = Font(bold=True, size=15, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=DARK)
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 30; row += 1
    # legend
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    c = ws.cell(row=row, column=1, value='*  =  Mandatory        ■  yellow  =  upload to folder')
    c.font = Font(size=9, italic=True, color=GREY)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[row].height = 16; row += 1
    # header
    for i, h in enumerate(['No.','Item / Field','Response','Upload → Folder / Remarks'], start=1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.font = Font(bold=True, size=10, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor=MID)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cell.border = border
    ws.row_dimensions[row].height = 18; row += 1

    key_rows = {}  # key -> 行号（用于 ref 跨模块引用）
    for m in MODULES:
        # 只有含上传项(upload)的模块才建附件文件夹
        has_upload = any(r['type'] == 'upload' for r in m['rows'])
        if m.get('folder') and has_upload:
            os.makedirs(os.path.join(out_dir, m['folder']), exist_ok=True)
        # module band
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        c = ws.cell(row=row, column=1, value=f"MODULE {m['id']}  ·  {m['title']}")
        c.font = Font(bold=True, size=13, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=DARK)
        c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        ws.row_dimensions[row].height = 26; row += 1

        fill_no = 0
        upload_no = 0
        group_number_start = None
        for r in m['rows']:
            t = r['type']

            if t == 'note':
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
                c = ws.cell(row=row, column=1, value=r['text'])
                if r.get('emph'):
                    c.font = Font(size=10, bold=True, color=REQ)
                    c.fill = PatternFill('solid', fgColor=UP)
                    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    ws.row_dimensions[row].height = 30
                else:
                    c.font = Font(size=8, italic=True, color=GREY)
                    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
                    ws.row_dimensions[row].height = 26
                row += 1
                continue

            if t == 'subhead':
                group_number_start = None
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
                c = ws.cell(row=row, column=1, value=r['text'])
                c.font = Font(bold=True, size=10, color='1F4E79')
                c.fill = PatternFill('solid', fgColor=SUBHEAD)
                c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
                ws.row_dimensions[row].height = 18; row += 1
                continue

            if t == 'fill':
                fill_no += 1
                req = r.get('required', False)
                zebra = PatternFill('solid', fgColor=LIGHT) if fill_no % 2 == 0 else None
                # No. 列：share_title 时用箭头占位（与前一行共享标题）
                if r.get('share_title'):
                    a = ws.cell(row=row, column=1, value='↳')
                else:
                    a = ws.cell(row=row, column=1, value=fill_no)
                a.font = Font(size=9, color='404040'); a.alignment = Alignment(horizontal='center', vertical='center')
                a.border = border
                if zebra: a.fill = zebra
                # B 列：share_title 时留空（共享上一行标题）
                if r.get('share_title'):
                    b = ws.cell(row=row, column=2, value='')
                else:
                    b = ws.cell(row=row, column=2, value=r['item'])
                b.font = Font(size=10, bold=req, color='1F1F1F' if req else GRAYTX)
                b.alignment = Alignment(vertical='center', wrap_text=True); b.border = border
                if zebra: b.fill = zebra
                # C 列
                d = ws.cell(row=row, column=3)
                if r.get('ref') and r['ref'] in key_rows:
                    d.value = f'=C{key_rows[r["ref"]]}'
                elif r.get('auto_sum') and group_number_start:
                    d.value = f'=SUM(C{group_number_start}:C{row-1})'
                else:
                    d.value = r.get('default','')
                    if r.get('ftype') == 'number' and group_number_start is None:
                        group_number_start = row
                if zebra: d.fill = zebra
                d.alignment = Alignment(vertical='top', wrap_text=True); d.border = border
                ftype = r.get('ftype','text')
                if ftype == 'dropdown':
                    opts = tuple(r.get('options', []))
                    if opts in opt_name:
                        dv = DataValidation(type='list',
                                            formula1=f'={opt_name[opts]}',
                                            allow_blank=True)
                        dv.showDropDown = None  # 不写该属性，Excel 默认显示下拉箭头
                        ws.add_data_validation(dv); dv.add(d)
                elif ftype == 'date':
                    d.number_format = 'YYYY-MM-DD'
                e = ws.cell(row=row, column=4)
                parts = []
                if req: parts.append('*Mandatory')
                if r.get('note'): parts.append(r['note'])
                e.value = ' · '.join(parts)
                e.font = Font(size=9, color=REQ if req else GRAYTX)
                e.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True); e.border = border
                if zebra: e.fill = zebra
                # 记录 key 行号
                if r.get('key'):
                    key_rows[r['key']] = row
                ws.row_dimensions[row].height = 24 if ftype == 'text_large' else 18
                row += 1
                continue

            if t == 'upload':
                upload_no += 1
                req = r.get('required', False)
                sub = f"{m['id']}{upload_no}"
                a = ws.cell(row=row, column=1, value=sub)
                a.font = Font(size=9, bold=True, color='7F6000'); a.alignment = Alignment(horizontal='center', vertical='center')
                a.fill = PatternFill('solid', fgColor=UP); a.border = border
                item_text = r['item']
                if r.get('template') == 'reference':
                    item_text += '  (reference sample)'
                elif r.get('template'):
                    item_text += '  (template available)'
                b = ws.cell(row=row, column=2, value=item_text)
                b.font = Font(size=10, bold=req, color='1F1F1F' if req else GRAYTX)
                b.alignment = Alignment(vertical='center', wrap_text=True)
                b.fill = PatternFill('solid', fgColor=UP); b.border = border
                c = ws.cell(row=row, column=3, value='See attachment')
                c.font = Font(size=9, italic=True, color='A6A6A6'); c.alignment = Alignment(vertical='center')
                c.fill = PatternFill('solid', fgColor=UP); c.border = border
                e = ws.cell(row=row, column=4)
                e.value = (('*Must upload  ·  ' if req else '') + f"→ {sub}  ({m['folder']})"
                           + (('\n' + r['note']) if r.get('note') else ''))
                e.font = Font(size=9, color=REQ if req else '7F6000')
                e.alignment = Alignment(vertical='center', wrap_text=True)
                e.fill = PatternFill('solid', fgColor=UP); e.border = border
                ws.row_dimensions[row].height = 34 if r.get('note') else 18
                row += 1
                continue

        row += 1  # gap between modules

    out_path = os.path.join(out_dir, 'ADNOC WISON SGP Subcontractor Registration.xlsx')
    wb.save(out_path)
    print('SAVED:', out_path)
    print('MODULES:', [m['id'] for m in MODULES])

if __name__ == '__main__':
    build()
