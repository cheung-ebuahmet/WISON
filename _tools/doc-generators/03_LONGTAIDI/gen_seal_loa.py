"""
Generate Seal/Chop Instruction for LONGTAIDI LOA.
"""
import sys, os
sys.stdout = open(sys.stdout.fileno(), mode='w', encoding='utf-8', errors='replace')

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

OUT_DIR = r"D:\Wison\Subcon_Payments\99 LONGTAIDI fob\Contract\Part I Subcontract Agreement\LOA"
OUT = os.path.join(OUT_DIR, "用印简要说明_LOA_LONGTAIDI.docx")

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)
style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

# ============================================================
# Helpers
# ============================================================
def ap(text, bold=False, size=10.5, align=None, color=None, sb=0, sa=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after = Pt(sa)
    if align: p.alignment = align
    run = p.add_run(text)
    run.font.name = 'Calibri'; run.font.size = Pt(size); run.bold = bold
    if color: run.font.color.rgb = color
    return p

def ah(text, size=12, sb=12, sa=4):
    return ap(text, bold=True, size=size, color=RGBColor(0x1F,0x4E,0x79), sb=sb, sa=sa)

def set_cell(cell, text, bold=False, size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    if p.runs and p.runs[0].text:
        p.runs[0].text = ''
    run = p.add_run(text) if not p.runs else p.add_run(text)
    run.font.name = 'Calibri'; run.font.size = Pt(size); run.bold = bold

def shd(cell, color):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>'))

def add_borders(table):
    tblPr = table._tbl.tblPr
    if tblPr is None:
        tblPr = parse_xml(f'<w:tblPr {nsdecls("w")}></w:tblPr>')
        table._tbl.insert(0, tblPr)
    b = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="888888"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="888888"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="888888"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="888888"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="888888"/>'
        '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="888888"/>'
        '</w:tblBorders>')
    tblPr.append(b)

# ============================================================
# TITLE
# ============================================================
ap("用印简要说明", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, sa=2,
   color=RGBColor(0x1F,0x4E,0x79))
ap("Seal / Chop Instruction", False, 9, WD_ALIGN_PARAGRAPH.CENTER, RGBColor(0x99,0x99,0x99), sa=10)

# ============================================================
# 一、用印文件
# ============================================================
ah("一、用印文件")

t1 = doc.add_table(rows=6, cols=2)
t1.alignment = WD_TABLE_ALIGNMENT.LEFT
add_borders(t1)
for i, w in enumerate([Inches(1.8), Inches(4.7)]):
    for row in t1.rows:
        row.cells[i].width = w

info = [
    ("文件名称", "《中标函》"),
    ("英文名称", "Letter of Award (LOA)\nfor Fire-Fighting Piping (3-inch and above) Fabrication Works on FOB Basis"),
    ("LOA 编号", "24108-CC0401-25002-LOA-19"),
    ("LOA 日期", "2026 年 7 月 28 日"),
    ("发函方 / 签约方", "WISON ENERGY ENGINEERING (HONG KONG) LIMITED – ABU DHABI\n（承包商 / Contractor —— 签署发出方）\n\n受函方：\nCangzhou Longtaidi Pipe Technology Co., Ltd.\n沧州隆泰迪管道科技有限公司\n（分包商 / Subcontractor —— 签署接受方）"),
    ("项目名称", "Sulphur Granulation Plant (RSGP) at RSHT-2 for Hail & Ghasha Project"),
]
for ri, (label, val) in enumerate(info):
    set_cell(t1.rows[ri].cells[0], label, bold=True)
    shd(t1.rows[ri].cells[0], 'F0F3F7')
    set_cell(t1.rows[ri].cells[1], val)

doc.add_paragraph()

# ============================================================
# 二、用印份数与接受方式
# ============================================================
ah("二、用印份数与接受方式")

ap("• LOA 由 Wison 签署盖章后发出，一式两份正本。", sa=2)
ap("• LONGTAIDI 收到后，在两份正本上签字盖章表示接受，其中一份返还 Wison 留存，一份 LONGTAIDI 自留。", sa=2)
ap("• 用印方式：公司公章（或合同专用章）+ 授权代表人签字。", sa=2)
ap("• 语言版本：英文为唯一权威语言。", sa=2)

doc.add_paragraph()

# ============================================================
# 三、LOA 核心内容
# ============================================================
ah("三、LOA 核心内容")

ap("本中标函为 Wison 向 LONGTAIDI 发出的正式授标文件，确认将如下工程授予 LONGTAIDI：", sa=4)

ap("1. 授标范围：消防管道（公称直径 3 寸及以上，法兰连接）预制加工工程，以 FOB 方式在中国指定港口交付。", size=10.5, sa=2)
ap("2. 暂定分包合同价格：USD 976,747.73（不含增值税），以实际验收工程量按全费用单价计量支付。适用中国增值税法规。", size=10.5, sa=2)
ap("3. 合同生效日 / 开工日：2026 年 7 月 28 日（LOA 签发之日即为合同生效日、开工日）。", size=10.5, sa=2)
ap("4. 履约保函：合同价格的 10%（USD 97,674.77），须在生效后 14 日内提交。", size=10.5, sa=2)
ap("5. 合同签署期限：双方须在 LOA 签发后 30 日内签署正式分包合同。签署后，LOA 自动失效并以正式分包合同为准。签署前，LOA 及其附件构成双方之间有约束力的协议。", size=10.5, sa=2)

ap("6. LOA 附件（共 4 个 Annex，为 LOA 不可分割的组成部分）：", size=10.5, sa=2)

t_annex = doc.add_table(rows=5, cols=3)
t_annex.alignment = WD_TABLE_ALIGNMENT.CENTER
add_borders(t_annex)
for i, w in enumerate([Inches(1.6), Inches(3.6), Inches(1.3)]):
    for row in t_annex.rows:
        row.cells[i].width = w

for ci, txt in enumerate(["附件编号", "附件名称", "状态"]):
    set_cell(t_annex.rows[0].cells[ci], txt, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    shd(t_annex.rows[0].cells[ci], '1F4E79')
    for run in t_annex.rows[0].cells[ci].paragraphs[0].runs:
        run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)

annexes = [
    ("Annex 1", "Key Milestone（关键里程碑）", "后续补充"),
    ("Annex 2", "Liquidated Damage（误期违约金）", "后续补充"),
    ("Annex 3", "No Deviation Declaration（无偏差声明）", "后续补充"),
    ("Annex 4", "ICV Improvement Plan（ICV 改进计划）", "N/A（FOB 中国交付不适用）"),
]
for ri, (no, name, status) in enumerate(annexes):
    r = ri + 1
    set_cell(t_annex.rows[r].cells[0], no, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(t_annex.rows[r].cells[1], name)
    set_cell(t_annex.rows[r].cells[2], status, align=WD_ALIGN_PARAGRAPH.CENTER)
    if ri % 2 == 0:
        for cell in t_annex.rows[r].cells:
            shd(cell, 'FAFBFC')

doc.add_paragraph()

# ============================================================
# 四、相关审批流程
# ============================================================
ah("四、相关审批流程")

ap("1. 合同评审表（Contract Review Form）—— 须随附上传，为用印前置必要条件。", sa=2)
ap("2. 本 LOA 为正式授标文件，须按公司授标审批权限完成内部批准后，方可用印发出。", sa=2)
ap("3. LONGTAIDI 方签署接受后，应扫描 LOA 正本 PDF 存档。项目管理部及商务部各留存一份扫描件。", sa=2)
ap("4. 后续正式分包合同（WISON24108C26010）签署后，本 LOA 自动失效，但作为合同文件组成部分（优先级第二，仅次于合同协议本身）继续具有参考效力。", sa=2)

doc.add_paragraph()

# ============================================================
# 五、随附上传附件清单
# ============================================================
ah("五、随附上传附件清单")

t2 = doc.add_table(rows=5, cols=4)
t2.alignment = WD_TABLE_ALIGNMENT.CENTER
add_borders(t2)

for i, w in enumerate([Inches(0.4), Inches(3.2), Inches(1.5), Inches(1.4)]):
    for row in t2.rows:
        row.cells[i].width = w

for ci, txt in enumerate(["序号", "附件名称", "文件大小参考", "备注"]):
    set_cell(t2.rows[0].cells[ci], txt, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    shd(t2.rows[0].cells[ci], '1F4E79')
    for run in t2.rows[0].cells[ci].paragraphs[0].runs:
        run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)

atts = [
    ("1", "合同评审表（已签批）", "~2 MB", "必须上传"),
    ("2", "LOA Main Body — LONGTAIDI\n（中标函正文，含签署页）", "~50 KB", "用印文件"),
    ("3", "LONGTAIDI 最终报价 / 商业建议书", "~1 MB", "授标依据，参考附后"),
    ("4", "LONGTAIDI 授权委托书（POA）/ 法定代表人证明", "~1 MB", "确认签字人权限"),
]
for ri, (no, name, size, remark) in enumerate(atts):
    r = ri + 1
    set_cell(t2.rows[r].cells[0], no, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(t2.rows[r].cells[1], name)
    set_cell(t2.rows[r].cells[2], size, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell(t2.rows[r].cells[3], remark, align=WD_ALIGN_PARAGRAPH.CENTER)
    if ri % 2 == 0:
        for cell in t2.rows[r].cells:
            shd(cell, 'FAFBFC')

doc.add_paragraph()

# ============================================================
# 六、备注
# ============================================================
ah("六、备注")

notes = [
    "本 LOA 受阿布扎比法律及阿联酋联邦法律管辖，争议解决机制与正式分包合同中的规定一致。",

    "LONGTAIDI 为中国公司，用印前须确认其签字人的授权委托书（POA）或法定代表人证明文件已提供并有效。如签字人非法定代表人，须附公司董事会决议或授权书。",

    "中标价格 USD 976,747.73 为暂定价格（不含增值税），最终结算以实际验收合格的工程量为准。",

    "Annex 1-3 目前占位为后续补充，不影响 LOA 的法律效力。Annex 4（ICV）已标记为 N/A。",

    "用印完成后，建议：LOA 正本两份均须有 Wison 盖章 + 签字，发出后 LONGTAIDI 签字盖章返还一份。返还件归档前请确认 LONGTAIDI 的签字盖章齐全、签字人身份与 POA 一致。",
]

for n in notes:
    ap("• " + n, size=10, sa=3)

# ============================================================
# END
# ============================================================
doc.add_paragraph()
ap("— End —", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0xAA,0xAA,0xAA))

doc.save(OUT)
print(f"Saved: {OUT}")
