import sys
sys.stdout.reconfigure(encoding='utf-8')

# File 1: PDF (already have text via pymupdf)
import fitz
print("=" * 60)
print("FILE 1: RSGP MEI Package Tender Briefing 2026-07-22.pdf")
print("=" * 60)
doc = fitz.open(r"D:\Wison\Project_Info\Tender Briefings\RSGP MEI Package Tender Briefing 2026-07-22.pdf")
print(f"Pages: {doc.page_count}")
for i, page in enumerate(doc):
    text = page.get_text()
    if text.strip():
        print(f"\n--- Page {i+1} ---")
        print(text[:2000])
doc.close()

# File 2: PPTX 2025-11-13
from pptx import Presentation
print("\n\n" + "=" * 60)
print("FILE 2: RSGP MEI Package Tender Briefing 2025-11-13.pptx")
print("=" * 60)
prs = Presentation(r"D:\Wison\Project_Info\Tender Briefings\RSGP MEI Package Tender Briefing 2025-11-13.pptx")
print(f"Slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i+1} ---")
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    texts.append(t)
    if texts:
        for t in texts:
            print(t[:500])
    else:
        print("(image/table only)")

# File 3: PPTX 2026-03-22
print("\n\n" + "=" * 60)
print("FILE 3: RSGP MEI Pkg I-3 Execution Strategy Meeting 2026-03-22.pptx")
print("=" * 60)
prs = Presentation(r"D:\Wison\Project_Info\Tender Briefings\RSGP MEI Pkg I-3 Execution Strategy Meeting 2026-03-22.pptx")
print(f"Slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    print(f"\n--- Slide {i+1} ---")
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    texts.append(t)
    if texts:
        for t in texts:
            print(t[:500])
    else:
        print("(image/table only)")
