"""Consolidate two CSVs into one dual-sheet Excel master index."""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SRC = r"D:\Wison\_Ref\kb"
MAIN = os.path.join(SRC, "00_Contract Master Index_6000066707_主合同文件主索引.csv")
SUMM = os.path.join(SRC, "00_Contract Master Index_6000066707_按ANNEXURE汇总.csv")
OUT  = os.path.join(SRC, "00_Contract Master Index_6000066707.xlsx")

# --- Styles ---
font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
font_body   = Font(name='Calibri', size=10)
fill_header = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
fill_gray   = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
fill_ocr    = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
fill_signed = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
fill_missing= PatternFill(start_color='F4CCCC', end_color='F4CCCC', fill_type='solid')
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9'),
)
align_wrap = Alignment(horizontal='left', vertical='center', wrap_text=True)

wb = Workbook()

# ========================================================================
# Sheet 1 — 汇总 Summary
# ========================================================================
ws1 = wb.active
ws1.title = "汇总 Summary"

# Read summary CSV
with open(SUMM, 'r', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    sum_rows = list(reader)

for ri, row in enumerate(sum_rows):
    for ci, val in enumerate(row):
        cell = ws1.cell(row=ri+1, column=ci+1, value=val)
        if ri == 0:
            cell.font = font_header
            cell.fill = fill_header
        else:
            cell.font = font_body
        cell.alignment = align_wrap
        cell.border = thin_border

# Column widths
ws1.column_dimensions['A'].width = 18
ws1.column_dimensions['B'].width = 38
ws1.column_dimensions['C'].width = 12
ws1.column_dimensions['D'].width = 10
ws1.column_dimensions['E'].width = 10
ws1.column_dimensions['F'].width = 10
ws1.column_dimensions['G'].width = 12

# Color-code summary rows by completeness
for ri in range(2, ws1.max_row + 1):
    pct_cell = ws1.cell(row=ri, column=7)
    try:
        pct = int(str(pct_cell.value).rstrip('%'))
    except (ValueError, AttributeError):
        continue
    for ci in range(1, ws1.max_column + 1):
        c = ws1.cell(row=ri, column=ci)
        if pct >= 90:
            c.fill = fill_ocr
        elif pct >= 50:
            c.fill = fill_signed
        else:
            c.fill = fill_missing

ws1.freeze_panes = 'A2'
ws1.auto_filter.ref = ws1.dimensions

# ========================================================================
# Sheet 2 — 主索引 Master Index
# ========================================================================
ws2 = wb.create_sheet("主索引 Master Index")

with open(MAIN, 'r', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    main_rows = list(reader)

# Row 1 = header, Row 2 = sub-header (English labels)
for ri, row in enumerate(main_rows):
    for ci, val in enumerate(row):
        cell = ws2.cell(row=ri+1, column=ci+1, value=val)
        if ri == 0:
            cell.font = font_header
            cell.fill = fill_header
        elif ri == 1:
            cell.font = Font(name='Calibri', size=9, italic=True, color='808080')
            cell.fill = fill_gray
        else:
            cell.font = font_body
            # Color rows by OCR status
            ocr_val = str(main_rows[ri][13]) if len(main_rows[ri]) > 13 else ''
            if 'OCR' in ocr_val or '✅' in ocr_val:
                for c2 in range(1, len(row) + 1):
                    ws2.cell(row=ri+1, column=c2).fill = fill_ocr
            elif 'Signed' in ocr_val or '📄' in ocr_val:
                for c2 in range(1, len(row) + 1):
                    ws2.cell(row=ri+1, column=c2).fill = fill_signed
        cell.alignment = align_wrap
        cell.border = thin_border

# Column widths — practical defaults
col_widths = {
    'A': 14,  # 层级
    'B': 7,   # 序号
    'C': 38,  # ANNEXURE
    'D': 42,  # EXHIBIT
    'E': 30,  # 文件名称
    'F': 12,  # 说明
    'G': 14,  # 负责部门
    'H': 10,  # 小签人
    'I': 7,   # REV
    'J': 12,  # Due Date
    'K': 8,   # 已核对
    'L': 8,   # 有变更
    'M': 10,  # 核对备注
    'N': 14,  # OCR状态
    'O': 50,  # 本地文件名
    'P': 55,  # 本地路径
}
for col_letter, width in col_widths.items():
    ws2.column_dimensions[col_letter].width = width

ws2.freeze_panes = 'A3'
ws2.auto_filter.ref = ws2.dimensions

wb.save(OUT)
print(f"Consolidated master index saved: {OUT}")

# Delete old CSVs
for old in [MAIN, SUMM]:
    os.remove(old)
    print(f"Deleted: {old}")
